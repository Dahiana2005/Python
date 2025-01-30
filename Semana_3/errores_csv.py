import pandas as pd
import logging
from pydantic import BaseModel, ValidationError, EmailStr, field_validator
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Configuración del logging
logging.basicConfig(filename='errores_validacion.log', level=logging.ERROR, format='%(asctime)s - %(levelname)s - %(message)s')

class Definicion_datos(BaseModel):
    Id_pedido: int
    Id_Cliente: str
    Email: EmailStr
    Telefono: str
    Prioridad: str

    @field_validator('Id_pedido')
    def validar_id_pedido(cls, value):
        if value < 0:
            raise ValueError('El id del pedido no puede ser negativo')
        return value

    @field_validator('Telefono')
    def validar_telefono(cls, value):
        if not value.isdigit() or len(value) != 10:
            raise ValueError('El número telefónico debe contener 10 caracteres numéricos')
        return value

    @field_validator('Prioridad')
    def validar_prioridad(cls, value):
        Lista_prioridades = {"Crítica", "Alta", "Baja", "Media"}
        if value not in Lista_prioridades:
            raise ValueError(f'La prioridad debe ser una de las siguientes: {Lista_prioridades}')
        return value

    @field_validator('Id_Cliente')
    def validar_id_cliente(cls, value):
        if not value:
            raise ValueError('ID_Cliente no puede estar vacío')
        return value

def resaltar_errores(nombre_archivo: str, filas_con_errores: list):
    df = pd.read_csv(nombre_archivo, encoding='utf-8', delimiter=';')
    df.to_excel('errores_resaltados.xlsx', index=False, engine='openpyxl')

    wb = load_workbook('errores_resaltados.xlsx')
    ws = wb.active

    rojo_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    for error in filas_con_errores:
        fila = error['Fila']
        
        for col in range(1, len(df.columns) + 1):
            ws.cell(row=fila, column=col).fill = rojo_fill

    wb.save('errores_resaltados.xlsx')

def validar_csv_pandas(nombre_archivo: str):
    try:
        df = pd.read_csv(nombre_archivo, encoding='utf-8', delimiter=';')
        print(df.columns)  # Mostrar las columnas del CSV
    except FileNotFoundError:
        print(f"Error: No se encontró el archivo '{nombre_archivo}'.")
        return
    except pd.errors.ParserError:
        print(f"Error al parsear el archivo CSV. Asegúrate de que el formato sea correcto.")
        return
    except Exception as e:
        print(f"Un error inesperado ocurrió al leer el archivo: {e}")
        return

    errores = []
    filas_con_errores = []

    for index, row in df.iterrows():
        if pd.notna(row['Telefono']):
            if isinstance(row['Telefono'], float): 
                telefono = str(int(row['Telefono']))
            else:
                telefono = str(row['Telefono']).replace('.0', '')
        else:
            telefono = ''

        data = {
            'Id_pedido': int(row['Id_pedido']),
            'Id_Cliente': row['Id_cliente'],
            'Email': row['Email'],
            'Telefono': telefono,
            'Prioridad': row['Prioridad']
        }

        try:
            Definicion_datos(**data)
        except ValueError as ve:
            errores.append({'Fila': index + 2, 'Error': str(ve)})
            filas_con_errores.append({'Fila': index + 2, 'Error': str(ve)})
            logging.error(f'Línea {index + 2} (Pandas index {index}): {ve}')
        except ValidationError as ve:
            for error in ve.errors():
                errores.append({'Fila': index + 2, 'Campo': error["loc"][0], 'Error': error["msg"]})
                filas_con_errores.append({'Fila': index + 2, 'Error': f"{error['loc'][0]} - {error['msg']}"})
                logging.error(f'Línea {index + 2} (Pandas index {index}): {error["loc"][0]} - {error["msg"]}')
        except Exception as e:
            errores.append({'Fila': index + 2, 'Error': f"Error inesperado: {e}"})
            filas_con_errores.append({'Fila': index + 2, 'Error': f"Error inesperado: {e}"})
            logging.error(f'Línea {index + 2} (Pandas index {index}): Error inesperado: {e}')

    if errores:
        print(f"Se encontraron {len(errores)} errores en el archivo CSV. Revise el archivo errores_validacion.log para más detalles.")

        resaltar_errores(nombre_archivo, filas_con_errores)
    else:
        print("No se encontraron errores en el archivo CSV")

if __name__ == "__main__":
    validar_csv_pandas('ventas_error.csv')
